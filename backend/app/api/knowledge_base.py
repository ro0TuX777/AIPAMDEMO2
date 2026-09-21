"""
Knowledge Base API endpoints for RAG — per-job scoped.

POST   /jobs/{job_id}/kb/documents          — upload a new KB document for a job (JSON text)
POST   /jobs/{job_id}/kb/upload-binary      — upload a binary file (PDF/DOCX/XLSX), extract text server-side
GET    /jobs/{job_id}/kb/documents          — list KB documents for a job
GET    /jobs/{job_id}/kb/documents/{doc_id} — get document detail (with content)
DELETE /jobs/{job_id}/kb/documents/{doc_id} — delete a document and its embeddings
POST   /jobs/{job_id}/kb/search             — search job's KB by query text
"""

from __future__ import annotations

import hashlib
import io
import logging
import uuid
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, Header, HTTPException, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.app.api.deps import get_db, verify_token
from backend.app.config_v2 import Settings, get_settings
from backend.app.models.job import Job
from backend.app.models.knowledge_base import KBDocument
from backend.app.schemas.knowledge_base import (
    KBDocumentCreate,
    KBDocumentDetail,
    KBDocumentListOut,
    KBDocumentOut,
)
from backend.app.services.kb_service import (
    EmbeddingError,
    GLOBAL_JOB_SENTINEL,
    delete_document as kb_delete_document,
    index_document as kb_index_document,
    retrieve as kb_retrieve,
)
from backend.app.services.embedding_models import get_embedding_model_service

logger = logging.getLogger("aipam.api.kb")

router = APIRouter(dependencies=[Depends(verify_token)], tags=["knowledge-base"])

def _now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%fZ")


def _vector_dir(settings: Settings) -> str:
    """Path to the ChromaDB vector store (sibling of the SQLite DB)."""
    return str(settings.aipam_db_path).replace("aipam.db", "vector_store")


def _embedding_config(settings: Settings):
    """Resolve the shared runtime model for this vector operation."""
    config = get_embedding_model_service(settings.aipam_ollama_url).get_active()
    if config is None:
        raise HTTPException(status_code=503, detail="No embedding model is selected")
    return config


def _require_job(db: Session, job_id: str) -> Job:
    """Validate the job exists, raise 404 otherwise."""
    job = db.get(Job, job_id)
    if not job:
        raise HTTPException(status_code=404, detail=f"Job {job_id} not found")
    return job


async def require_kb_admin(
    x_kb_admin_token: str | None = Header(None, alias="X-KB-Admin-Token"),
    settings: Settings = Depends(get_settings),
) -> None:
    """Gate for global-library mutations (create / delete / re-index).

    When ``aipam_kb_admin_token`` is configured, these require that token via the
    X-KB-Admin-Token header — on top of the normal API bearer — so shared
    material can only be curated by admins. Unset (default) → open, so this is
    backward-compatible. Reads and chat retrieval are never gated (the library
    exists to ground the AI for everyone).
    """
    admin_token = (settings.aipam_kb_admin_token or "").strip()
    if not admin_token:
        return
    if (x_kb_admin_token or "") != admin_token:
        raise HTTPException(
            status_code=403,
            detail="Managing the shared library requires the KB admin token.",
        )


DEGRADED_MSG = (
    "Embedding model unavailable — indexed with low-quality hash fallbacks, so "
    "semantic search will be poor. Re-index once the embedding model is pulled."
)


def _create_pending_doc(
    db: Session,
    *,
    job_id: str | None,
    name: str,
    doc_type: str,
    description: str | None,
    content: str,
    filename: str | None = None,
) -> KBDocument:
    """Persist a KB document as ``pending``. Indexing happens in the background.

    ``job_id=None`` creates a global reference-library document. Rejects (409) an
    upload whose content is byte-for-byte identical to an existing document in
    the same scope, so re-uploading the same manual doesn't duplicate chunks.
    """
    content_sha = hashlib.sha256(content.encode("utf-8")).hexdigest()
    scope = KBDocument.job_id.is_(None) if job_id is None else (KBDocument.job_id == job_id)
    existing = db.execute(
        select(KBDocument).where(KBDocument.content_sha256 == content_sha, scope)
    ).scalars().first()
    if existing is not None:
        raise HTTPException(
            status_code=409,
            detail=f"This document is already in the knowledge base as '{existing.name}'.",
        )

    now = _now_iso()
    doc = KBDocument(
        id=str(uuid.uuid4()),
        job_id=job_id,
        name=name,
        doc_type=doc_type,
        description=description or None,
        filename=filename,
        content=content,
        content_sha256=content_sha,
        status="pending",
        created_at=now,
        updated_at=now,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


async def _index_doc_task(
    doc_id: str,
    bind: Any,
    ollama_url: str,
    persist_dir: str,
    delete_first: bool = False,
) -> None:
    """Background task: chunk, embed, and store a document, then update status.

    Runs after the HTTP response so large manuals don't block the request or
    time it out. Uses its own DB session bound to the same engine/connection as
    the request that scheduled it.
    """
    db = Session(bind=bind)
    try:
        doc = db.get(KBDocument, doc_id)
        if doc is None:
            return
        try:
            embedding_config = get_embedding_model_service(ollama_url).get_active()
            if embedding_config is None:
                raise EmbeddingError("No embedding model is selected")
            if delete_first:
                await kb_delete_document(
                    doc_id=doc_id,
                    embedding_config=embedding_config,
                    persist_dir=persist_dir,
                )
            result = await kb_index_document(
                doc_id=doc_id,
                content=doc.content,
                doc_name=doc.name,
                doc_type=doc.doc_type,
                job_id=doc.job_id or GLOBAL_JOB_SENTINEL,
                ollama_url=ollama_url,
                embedding_config=embedding_config,
                persist_dir=persist_dir,
            )
            doc.chunk_count = result.chunk_count
            doc.status = "indexed"
            doc.error_message = None
        except Exception as exc:
            logger.exception("Failed to index KB document %s", doc_id)
            doc.status = "error"
            doc.error_message = str(exc)[:500]
        doc.updated_at = _now_iso()
        db.commit()
    finally:
        db.close()


def _schedule_index(
    background: BackgroundTasks,
    db: Session,
    settings: Settings,
    doc_id: str,
    *,
    delete_first: bool = False,
) -> None:
    """Queue background (re)indexing for a document."""
    background.add_task(
        _index_doc_task,
        doc_id,
        db.get_bind(),
        settings.aipam_ollama_url.rstrip("/"),
        _vector_dir(settings),
        delete_first,
    )


@router.post("/jobs/{job_id}/kb/documents", response_model=KBDocumentOut, status_code=201)
async def create_document(
    job_id: str,
    body: KBDocumentCreate,
    background: BackgroundTasks,
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Add a KB document for a job. Returns immediately; indexing runs in the background."""
    _require_job(db, job_id)
    doc = _create_pending_doc(
        db,
        job_id=job_id,
        name=body.name,
        doc_type=body.doc_type,
        description=body.description,
        content=body.content,
    )
    _schedule_index(background, db, settings, doc.id)
    return KBDocumentOut.model_validate(doc)


@router.get("/jobs/{job_id}/kb/documents", response_model=KBDocumentListOut)
async def list_documents(
    job_id: str,
    doc_type: str | None = None,
    db: Session = Depends(get_db),
):
    """List knowledge base documents for a specific job."""
    _require_job(db, job_id)
    query = select(KBDocument).where(KBDocument.job_id == job_id).order_by(KBDocument.created_at.desc())
    if doc_type:
        query = query.where(KBDocument.doc_type == doc_type)

    docs = db.execute(query).scalars().all()
    return KBDocumentListOut(
        items=[KBDocumentOut.model_validate(d) for d in docs],
        total=len(docs),
    )


@router.get("/jobs/{job_id}/kb/documents/{doc_id}", response_model=KBDocumentDetail)
async def get_document(
    job_id: str,
    doc_id: str,
    db: Session = Depends(get_db),
):
    """Get a knowledge base document with full content."""
    _require_job(db, job_id)
    doc = db.get(KBDocument, doc_id)
    if not doc or doc.job_id != job_id:
        raise HTTPException(status_code=404, detail="Document not found")
    return KBDocumentDetail.model_validate(doc)


@router.delete("/jobs/{job_id}/kb/documents/{doc_id}", status_code=204)
async def delete_document(
    job_id: str,
    doc_id: str,
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Delete a knowledge base document and its vector embeddings."""
    _require_job(db, job_id)
    doc = db.get(KBDocument, doc_id)
    if not doc or doc.job_id != job_id:
        raise HTTPException(status_code=404, detail="Document not found")

    # Delete from ChromaDB
    try:
        await kb_delete_document(
            doc_id=doc_id,
            embedding_config=_embedding_config(settings),
            persist_dir=_vector_dir(settings),
        )
    except Exception as exc:
        logger.warning("Failed to delete KB embeddings for %s: %s", doc_id, exc)

    db.delete(doc)
    db.commit()
    return None


@router.post("/jobs/{job_id}/kb/documents/{doc_id}/reindex", response_model=KBDocumentOut)
async def reindex_document(
    job_id: str,
    doc_id: str,
    background: BackgroundTasks,
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Re-chunk and re-embed a document (e.g. to recover from a degraded index)."""
    _require_job(db, job_id)
    doc = db.get(KBDocument, doc_id)
    if not doc or doc.job_id != job_id:
        raise HTTPException(status_code=404, detail="Document not found")
    doc.status = "pending"
    doc.error_message = None
    doc.updated_at = _now_iso()
    db.commit()
    db.refresh(doc)
    _schedule_index(background, db, settings, doc_id, delete_first=True)
    return KBDocumentOut.model_validate(doc)


BINARY_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls", ".pptx"}
MAX_BINARY_SIZE = 10 * 1024 * 1024  # 10 MB for binary files
_OCR_PAGE_CAP = 40  # OCR is slow (synchronous) — bound how many pages we render


def _ocr_pdf_pages(data: bytes, page_indexes: list[int]) -> dict[int, str]:
    """OCR specific PDF pages (0-based). Returns {index: text}.

    Best-effort: needs pytesseract + pdf2image (and the tesseract/poppler system
    binaries). Returns {} if that stack is unavailable, so an image-only PDF
    still degrades to a clear 422 rather than a 500.
    """
    if not page_indexes:
        return {}
    try:
        import pytesseract
        from pdf2image import convert_from_bytes
    except ImportError:
        logger.info("OCR stack (pytesseract/pdf2image) not installed — skipping OCR")
        return {}

    out: dict[int, str] = {}
    for idx in page_indexes[:_OCR_PAGE_CAP]:
        try:
            images = convert_from_bytes(data, dpi=200, first_page=idx + 1, last_page=idx + 1)
            if images:
                text = pytesseract.image_to_string(images[0]) or ""
                if text.strip():
                    out[idx] = text.strip()
        except Exception as exc:
            logger.warning("OCR failed for PDF page %d: %s", idx + 1, exc)
    if len(page_indexes) > _OCR_PAGE_CAP:
        logger.warning("OCR capped at %d pages (%d image pages present)", _OCR_PAGE_CAP, len(page_indexes))
    return out


def _extract_text_from_binary(filename: str, data: bytes) -> str:
    """Extract plain text from a binary document (PDF, DOCX, XLSX, PPTX)."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == ".pdf":
        try:
            import pdfplumber
        except ImportError:
            raise HTTPException(status_code=500, detail="pdfplumber is not installed on the server")
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            n_pages = len(pdf.pages)
            text_by_page: dict[int, str] = {}
            empty_pages: list[int] = []
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if text.strip():
                    text_by_page[i] = text
                else:
                    empty_pages.append(i)

        # OCR the pages that had no embedded text (scanned pages / screenshots).
        ocr_by_page = _ocr_pdf_pages(data, empty_pages)

        pages: list[str] = []
        for i in range(n_pages):
            if i in text_by_page:
                pages.append(f"--- Page {i + 1} ---\n{text_by_page[i]}")
            elif i in ocr_by_page:
                pages.append(f"--- Page {i + 1} (OCR) ---\n{ocr_by_page[i]}")
        if not pages:
            raise HTTPException(
                status_code=422,
                detail="PDF contains no extractable text — it looks image-only "
                "(scanned/screenshots) and OCR could not read it or is not "
                "installed. Paste the text or upload a text-based version.",
            )
        return "\n\n".join(pages)

    elif ext == ".pptx":
        try:
            from pptx import Presentation
        except ImportError:
            raise HTTPException(status_code=500, detail="python-pptx is not installed on the server")
        prs = Presentation(io.BytesIO(data))
        slides_text = []
        for i, slide in enumerate(prs.slides):
            lines: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = "".join(run.text for run in para.runs).strip()
                        if t:
                            lines.append(t)
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            lines.append("\t".join(cells))
            if lines:
                slides_text.append(f"--- Slide {i + 1} ---\n" + "\n".join(lines))
        if not slides_text:
            raise HTTPException(status_code=422, detail="Presentation contains no extractable text")
        return "\n\n".join(slides_text)

    elif ext == ".docx":
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise HTTPException(status_code=500, detail="python-docx is not installed on the server")
        doc = DocxDocument(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append("\t".join(cells))
        if not paragraphs:
            raise HTTPException(status_code=422, detail="DOCX contains no extractable text")
        return "\n".join(paragraphs)

    elif ext in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(status_code=500, detail="openpyxl is not installed on the server")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    rows.append("\t".join(cells))
            if rows:
                sheets_text.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
        wb.close()
        if not sheets_text:
            raise HTTPException(status_code=422, detail="Excel file contains no data")
        return "\n\n".join(sheets_text)

    else:
        raise HTTPException(status_code=400, detail=f"Unsupported binary file type: {ext}")


@router.post("/jobs/{job_id}/kb/upload-binary", response_model=KBDocumentOut, status_code=201)
async def upload_binary_document(
    job_id: str,
    background: BackgroundTasks,
    file: UploadFile = File(...),
    name: str = Form(...),
    doc_type: str = Form("other"),
    description: str = Form(""),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Upload a binary file (PDF, DOCX, XLSX) for a job; extract text and index in the background."""
    _require_job(db, job_id)

    filename = file.filename or "document"
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in BINARY_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}. Accepted: {', '.join(sorted(BINARY_EXTENSIONS))}")

    data = await file.read()
    if len(data) > MAX_BINARY_SIZE:
        raise HTTPException(status_code=413, detail=f"File too large (max {MAX_BINARY_SIZE // (1024*1024)}MB)")

    # Extract text
    content = _extract_text_from_binary(filename, data)
    if len(content) > 5_000_000:
        raise HTTPException(status_code=413, detail=f"Extracted text too large ({len(content):,} chars, max 5,000,000)")

    doc = _create_pending_doc(
        db,
        job_id=job_id,
        name=name,
        doc_type=doc_type,
        description=description or None,
        content=content,
        filename=filename,
    )
    _schedule_index(background, db, settings, doc.id)
    return KBDocumentOut.model_validate(doc)


class KBSearchRequest(BaseModel):
    query: str = Field(..., min_length=1, max_length=2000)
    n_results: int = Field(default=5, ge=1, le=20)
    doc_type: str | None = None


class KBSearchResult(BaseModel):
    text: str
    doc_name: str
    doc_type: str
    doc_id: str
    section: str = ""
    score: float


class KBSearchResponse(BaseModel):
    results: list[KBSearchResult]
    query: str


@router.post("/jobs/{job_id}/kb/search", response_model=KBSearchResponse)
async def search_kb(
    job_id: str,
    body: KBSearchRequest,
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Search a job's knowledge base (its own docs ∪ the global library)."""
    _require_job(db, job_id)

    results = await kb_retrieve(
        query=body.query,
        n_results=body.n_results,
        doc_type_filter=body.doc_type,
        job_id=job_id,
        ollama_url=settings.aipam_ollama_url.rstrip("/"),
        embedding_config=_embedding_config(settings),
        persist_dir=_vector_dir(settings),
    )

    return KBSearchResponse(
        results=[KBSearchResult(**r) for r in results],
        query=body.query,
    )


# ── Global Reference Library (job-less) ──────────────────────────────────
# Documents here are available to every analysis. Ideal for exploit user
# guides, capability manuals, SOC playbooks, and other durable references that
# should not be re-uploaded per job.


@router.get("/kb/library/config")
async def library_config(settings: Settings = Depends(get_settings)):
    """Whether curating the shared library requires the KB admin token."""
    return {"admin_required": bool((settings.aipam_kb_admin_token or "").strip())}


@router.post("/kb/library/documents", response_model=KBDocumentOut, status_code=201,
             dependencies=[Depends(require_kb_admin)])
async def create_library_document(
    body: KBDocumentCreate,
    background: BackgroundTasks,
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Add a document to the global reference library; indexing runs in the background."""
    doc = _create_pending_doc(
        db,
        job_id=None,
        name=body.name,
        doc_type=body.doc_type,
        description=body.description,
        content=body.content,
    )
    _schedule_index(background, db, settings, doc.id)
    return KBDocumentOut.model_validate(doc)


@router.post("/kb/library/upload-binary", response_model=KBDocumentOut, status_code=201,
             dependencies=[Depends(require_kb_admin)])
async def upload_library_binary_document(
    background: BackgroundTasks,
    file: UploadFile = File(...),
    name: str = Form(...),
    doc_type: str = Form("other"),
    description: str = Form(""),
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Upload a binary file (PDF, DOCX, XLSX) to the global reference library."""
    filename = file.filename or "document"
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in BINARY_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}. Accepted: {', '.join(sorted(BINARY_EXTENSIONS))}")

    data = await file.read()
    if len(data) > MAX_BINARY_SIZE:
        raise HTTPException(status_code=413, detail=f"File too large (max {MAX_BINARY_SIZE // (1024*1024)}MB)")

    content = _extract_text_from_binary(filename, data)
    if len(content) > 5_000_000:
        raise HTTPException(status_code=413, detail=f"Extracted text too large ({len(content):,} chars, max 5,000,000)")

    doc = _create_pending_doc(
        db,
        job_id=None,
        name=name,
        doc_type=doc_type,
        description=description or None,
        content=content,
        filename=filename,
    )
    _schedule_index(background, db, settings, doc.id)
    return KBDocumentOut.model_validate(doc)


@router.get("/kb/library/documents", response_model=KBDocumentListOut)
async def list_library_documents(
    doc_type: str | None = None,
    db: Session = Depends(get_db),
):
    """List all global reference-library documents."""
    query = select(KBDocument).where(KBDocument.job_id.is_(None)).order_by(KBDocument.created_at.desc())
    if doc_type:
        query = query.where(KBDocument.doc_type == doc_type)
    docs = db.execute(query).scalars().all()
    return KBDocumentListOut(
        items=[KBDocumentOut.model_validate(d) for d in docs],
        total=len(docs),
    )


@router.get("/kb/library/documents/{doc_id}", response_model=KBDocumentDetail)
async def get_library_document(
    doc_id: str,
    db: Session = Depends(get_db),
):
    """Get a global reference-library document with full content."""
    doc = db.get(KBDocument, doc_id)
    if not doc or doc.job_id is not None:
        raise HTTPException(status_code=404, detail="Document not found")
    return KBDocumentDetail.model_validate(doc)


@router.delete("/kb/library/documents/{doc_id}", status_code=204,
               dependencies=[Depends(require_kb_admin)])
async def delete_library_document(
    doc_id: str,
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Delete a global reference-library document and its embeddings."""
    doc = db.get(KBDocument, doc_id)
    if not doc or doc.job_id is not None:
        raise HTTPException(status_code=404, detail="Document not found")

    try:
        await kb_delete_document(
            doc_id=doc_id,
            embedding_config=_embedding_config(settings),
            persist_dir=_vector_dir(settings),
        )
    except Exception as exc:
        logger.warning("Failed to delete KB embeddings for %s: %s", doc_id, exc)

    db.delete(doc)
    db.commit()
    return None


@router.post("/kb/library/documents/{doc_id}/reindex", response_model=KBDocumentOut,
             dependencies=[Depends(require_kb_admin)])
async def reindex_library_document(
    doc_id: str,
    background: BackgroundTasks,
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Re-chunk and re-embed a global reference-library document."""
    doc = db.get(KBDocument, doc_id)
    if not doc or doc.job_id is not None:
        raise HTTPException(status_code=404, detail="Document not found")
    doc.status = "pending"
    doc.error_message = None
    doc.updated_at = _now_iso()
    db.commit()
    db.refresh(doc)
    _schedule_index(background, db, settings, doc_id, delete_first=True)
    return KBDocumentOut.model_validate(doc)


@router.post("/kb/library/search", response_model=KBSearchResponse)
async def search_library(
    body: KBSearchRequest,
    db: Session = Depends(get_db),
    settings: Settings = Depends(get_settings),
):
    """Search only the global reference library by semantic similarity."""
    results = await kb_retrieve(
        query=body.query,
        n_results=body.n_results,
        doc_type_filter=body.doc_type,
        job_id=None,
        include_global=True,
        ollama_url=settings.aipam_ollama_url.rstrip("/"),
        embedding_config=_embedding_config(settings),
        persist_dir=_vector_dir(settings),
    )
    return KBSearchResponse(
        results=[KBSearchResult(**r) for r in results],
        query=body.query,
    )
